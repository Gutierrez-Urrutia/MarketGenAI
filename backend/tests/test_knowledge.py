"""Tests de la base de conocimiento RAG (RF-VTA-01)."""
from __future__ import annotations

import io
from unittest.mock import AsyncMock, patch

import pytest

from app.services import rag_proposal_service as rag
from app.services.document_text_service import DocumentExtractionError, extract_text
from tests.conftest import FAKE_USER_SUB

API = "/api/v1/knowledge/documents"


def _docx_bytes() -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("Propuesta Contact Center", 1)
    d.add_paragraph("Servicio BPO de atención al cliente 24x7.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Servicio", "Precio"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Mesa de ayuda", "1200"
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Tarifas"
    ws.append(["Servicio", "Precio"]); ws.append(["Cobranza", 800])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _pdf_bytes() -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO(); c = canvas.Canvas(buf); c.drawString(72, 700, "Propuesta de outsourcing contable"); c.save()
    return buf.getvalue()


def test_extract_docx_includes_tables():
    text, ext = extract_text("p.docx", _docx_bytes())
    assert ext == ".docx" and "Contact Center" in text and "Mesa de ayuda" in text and "Precio: 1200" in text


def test_extract_xlsx_and_csv():
    assert "Servicio: Cobranza | Precio: 800" in extract_text("t.xlsx", _xlsx_bytes())[0]
    assert "Servicio: Soporte | Precio: 5" in extract_text("t.csv", "Servicio;Precio\nSoporte;5\n".encode())[0]


def test_extract_pdf():
    assert "outsourcing contable" in extract_text("p.pdf", _pdf_bytes())[0]


@pytest.mark.parametrize("name,data,msg", [
    ("a.doc", b"x", ".docx"),
    ("a.exe", b"x", "no soportado"),
    ("a.txt", b"", "vacío"),
    ("a.pdf", b"no es pdf", "PDF"),
])
def test_extract_rejects_invalid(name, data, msg):
    with pytest.raises(DocumentExtractionError, match=msg):
        extract_text(name, data)


def test_chunk_text_respects_size_and_splits_long_paragraphs():
    chunks = rag.chunk_text("palabra " * 1000)
    assert len(chunks) > 1 and all(len(c) <= rag.CHUNK_SIZE + rag.CHUNK_OVERLAP + 2 for c in chunks)


def _chunk(doc_id, text):
    return {"docId": doc_id, "text": text, "tokens": rag._tokens(text)}


def test_rank_prefers_matching_service_and_client_and_drops_irrelevant():
    chunks = [
        _chunk("a", "Servicio de cobranza telefónica para banca con gestión de mora"),
        _chunk("b", "Servicio de mesa de ayuda nivel 1 y 2 con SLAs mensuales"),
        _chunk("c", "Receta de cocina con harina"),
    ]
    docs = {"a": {"title": "Cobranza", "client": "Banco Uno"}, "b": {"title": "Help Desk"}, "c": {}}
    top = rag.rank_chunks("cobranza para banco", chunks, docs, client_name="Banco Uno")
    assert top[0]["docId"] == "a"
    assert all(c["docId"] != "c" for c in top)
    assert rag.rank_chunks("zzz inexistente", chunks, docs) == []


@pytest.mark.asyncio
async def test_get_rag_context_empty_without_documents():
    with patch.object(rag.knowledge_documents_repo, "list", new_callable=AsyncMock, return_value=[]):
        assert await rag.get_rag_context(FAKE_USER_SUB, "Cliente", "BPO") == ""
    assert await rag.get_rag_context(None, "Cliente", "BPO") == ""


@pytest.mark.asyncio
async def test_get_rag_context_uses_user_documents():
    docs = [{"id": "a", "title": "Propuesta Cobranza", "service": "cobranza", "client": ""}]
    chunks = [{"docId": "a", "text": "Cobranza BPO con 50 agentes", "tokens": rag._tokens("Cobranza BPO con 50 agentes")}]
    with (
        patch.object(rag.knowledge_documents_repo, "list", new_callable=AsyncMock, return_value=docs),
        patch.object(rag.knowledge_chunks_repo, "list", new_callable=AsyncMock, return_value=chunks) as lc,
    ):
        ctx = await rag.get_rag_context(FAKE_USER_SUB, "Acme", "servicio de cobranza")
    assert "Propuesta Cobranza" in ctx and "50 agentes" in ctx
    assert lc.await_args.kwargs["filters"] == [("userId", "==", FAKE_USER_SUB)]


@pytest.fixture
def manager(app):
    from app.dependencies.auth import CurrentUser, get_current_user
    original = app.dependency_overrides.get(get_current_user)
    base = original()  # type: ignore[misc]
    user = CurrentUser(**{**(base.model_dump() if hasattr(base, "model_dump") else vars(base)), "roles": ["manager"]})
    app.dependency_overrides[get_current_user] = lambda: user
    yield
    app.dependency_overrides[get_current_user] = original


@pytest.mark.asyncio
async def test_upload_requires_manager(client):
    resp = await client.post(API, files=[("files", ("a.txt", b"hola", "text/plain"))])
    assert resp.status_code == 403


@pytest.mark.asyncio
async def test_upload_partial_success(client, manager):
    fake_doc = {"id": "d1", "title": "ok.docx", "filename": "ok.docx", "fileType": "docx"}
    with patch.object(rag.knowledge_documents_repo, "list", new_callable=AsyncMock, return_value=[]), \
         patch.object(rag.knowledge_documents_repo, "create", new_callable=AsyncMock, return_value=fake_doc), \
         patch.object(rag.knowledge_chunks_repo, "create", new_callable=AsyncMock):
        resp = await client.post(API, files=[
            ("files", ("ok.docx", _docx_bytes(), "application/octet-stream")),
            ("files", ("malo.exe", b"x", "application/octet-stream")),
        ])
    body = resp.json()
    assert resp.status_code == 201
    assert [d["id"] for d in body["uploaded"]] == ["d1"] and body["failed"][0]["filename"] == "malo.exe"


@pytest.mark.asyncio
async def test_delete_other_users_document_forbidden(client, manager):
    with patch("app.routers.knowledge.knowledge_documents_repo.get", new_callable=AsyncMock,
               return_value={"id": "d1", "userId": "otro"}):
        assert (await client.delete(f"{API}/d1")).status_code == 403
