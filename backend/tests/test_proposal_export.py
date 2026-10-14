"""Exportación de propuestas a PDF/DOCX: texto real, idioma, total solo bajo la tabla de precios."""
from __future__ import annotations

from unittest.mock import AsyncMock, patch

import pytest
from pypdf import PdfReader

from app.routers import proposals as pr
from app.workers.tasks.proposal_tasks import build_pricing_table_html
from tests.conftest import FAKE_USER_SUB

API = "/api/v1/proposals"

TEAM_TABLE = "<table><thead><tr><th>Rol</th><th>Cantidad</th></tr></thead><tbody><tr><td>Gestor</td><td>15</td></tr></tbody></table>"


def _bpo_html() -> str:
    pricing = build_pricing_table_html(
        [{"name": "Gestión de cobranza", "description": "Gestor/mes", "qty": 15, "unitPrice": 1250}], 18750, "es"
    )
    return (
        "<h1>Propuesta de Cobranza Telefónica</h1><p>Para Financiera Horizonte</p>"
        "<h2>Resumen Ejecutivo</h2><p>Recuperación de cartera morosa con gestores dedicados.</p>"
        f"<h2>Equipo y Organización del Trabajo</h2>{TEAM_TABLE}"
        f"<h2>Precios y Condiciones Comerciales</h2>{pricing}<ul><li>Facturación mensual</li></ul>"
        "<h2>Condiciones Generales</h2><ul><li>Confidencialidad</li></ul>"
    )


def _proposal(**over) -> dict:
    base = {
        "id": "p1", "userId": FAKE_USER_SUB, "title": "Propuesta", "clientName": "Financiera Horizonte",
        "language": "es", "totalAmount": 18750, "content": _bpo_html(),
        "template": "NoonDalton Standard Proposal",
        "structuredContent": {"contact": "María Fuentes", "client": "Financiera Horizonte",
                              "industry": "B2B Operations", "serviceLine": "NoonDalton Standard Proposal"},
    }
    base.update(over)
    return base


def _pdf_text(path) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(str(path)).pages)


def test_pdf_has_real_selectable_text_in_spanish(tmp_path):
    proposal = _proposal()
    out = tmp_path / "p.pdf"
    pr._build_proposal_pdf(out, "Propuesta", proposal["content"], pr._proposal_cover_fields(proposal))
    text = _pdf_text(out)
    assert "Propuesta de Cobranza Telefónica" in text
    assert "Resumen Ejecutivo" in text and "Condiciones Generales" in text
    assert "FECHA DE EMISIÓN" in text and "VÁLIDA HASTA" in text
    assert "Página 1" in text
    assert "Aceptación" in text and "Firma" in text
    assert "NoonDalton Standard Proposal" not in text  # el nombre de plantilla no es dato del cliente
    assert "B2B Operations" not in text
    assert "AI Marketing Suite" not in text  # nombre interno del producto
    assert "Documento confidencial" in text
    assert "$18.750 USD" in text


def test_total_card_only_after_pricing_table(tmp_path):
    proposal = _proposal()
    out = tmp_path / "p.pdf"
    pr._build_proposal_pdf(out, "Propuesta", proposal["content"], pr._proposal_cover_fields(proposal))
    text = _pdf_text(out)
    assert text.count("Total de la propuesta") == 1  # la tabla de equipo no lleva total
    assert text.index("Gestor") < text.index("Total de la propuesta")


def test_legacy_table_without_class_still_gets_total(tmp_path):
    proposal = _proposal(content="<h2>Costos</h2>" + TEAM_TABLE)
    out = tmp_path / "p.pdf"
    pr._build_proposal_pdf(out, "P", proposal["content"], pr._proposal_cover_fields(proposal))
    assert "Total de la propuesta" in _pdf_text(out)


def test_split_accepts_frontend_presentation_html():
    wrapped = (
        '<article class="proposal-preview-document"><div class="proposal-paper">'
        '<header class="proposal-cover"><h1>Mi Propuesta</h1><p class="proposal-subtitle">Para ACME</p></header>'
        "<section><h2>Resumen Ejecutivo</h2><p>Texto.</p></section>"
        '<section><h2>Riesgos</h2><div class="proposal-card-list">'
        '<div class="proposal-detail-card"><strong>Mora alta</strong><p>Impacta caja.</p></div></div></section>'
        '<div class="proposal-total-card"><span>Total</span></div>'
        '<section class="proposal-acceptance"><h2>Aceptación</h2></section>'
        "</div></article>"
    )
    title, subtitle, tags = pr._split_proposal_content(wrapped)
    assert title == "Mi Propuesta" and subtitle == "Para ACME"
    texts = [t.get_text(" ", strip=True) for t in tags]
    assert texts == ["Resumen Ejecutivo", "Texto.", "Riesgos", "Mora alta: Impacta caja."]


def test_docx_export_localized(tmp_path):
    from docx import Document
    proposal = _proposal(language="pt")
    out = tmp_path / "p.docx"
    pr._build_proposal_docx(out, "Proposta", proposal["content"], pr._proposal_cover_fields(proposal))
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "DATA DE EMISSÃO" in text and "Aceitação" in text and "Total da proposta" in text
    footer = "\n".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "Documento confidencial" in footer and "Página" in footer
    assert "PAGE" in doc.sections[0].footer._element.xml  # campo de número de página


@pytest.mark.asyncio
async def test_export_endpoint_uses_edited_content_and_returns_pdf(client, tmp_path):
    edited = "<h1>Título Editado</h1><h2>Resumen Ejecutivo</h2><p>Texto editado por el usuario.</p>"
    with (
        patch("app.routers.proposals.proposals_repo.get", new_callable=AsyncMock, return_value=_proposal()),
        patch("app.routers.proposals._output_dir", return_value=tmp_path),
    ):
        resp = await client.post(f"{API}/p1/export", json={"format": "pdf", "content": edited, "issuedBy": "Admin User"})
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/pdf"
    out = tmp_path / "out.pdf"
    out.write_bytes(resp.content)
    text = _pdf_text(out)
    assert "Texto editado por el usuario." in text and "Admin User" in text


@pytest.mark.asyncio
async def test_export_endpoint_rejects_other_users_proposal(client):
    with patch("app.routers.proposals.proposals_repo.get", new_callable=AsyncMock, return_value=_proposal(userId="otro")):
        resp = await client.post(f"{API}/p1/export", json={"format": "pdf"})
    assert resp.status_code == 403
