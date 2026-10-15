"""
Genera 3 propuestas BPO ficticias en Word para probar la base de conocimiento RAG (RF-VTA-01).

Uso (desde backend/):
    .venv/Scripts/python.exe scripts/gen_propuestas_bpo_ejemplo.py tests/fixtures/propuestas_bpo
Los clientes, cifras y precios son inventados.
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

EMPRESA = (
    "NoonDalton BPO Services es una empresa de outsourcing de procesos de negocio con operación 24x7, "
    "centros de atención en Chile y Colombia y más de 12 años de experiencia. Su diferenciador son las "
    "economías de escala, la gestión por SLAs y la transparencia mediante reportes mensuales."
)

PROPUESTAS = [
    dict(
        fname="Propuesta_BPO_Contact_Center_Banco_Andino.docx",
        title="Propuesta de Contact Center Omnicanal",
        client="Banco Andino S.A.",
        resumen=(
            "Banco Andino enfrenta un alza sostenida de consultas en canales telefónico, correo y chat, con un "
            "nivel de servicio del 62% y un abandono de llamadas cercano al 14%. NoonDalton propone un contact "
            "center omnicanal de atención al cliente con 40 agentes, supervisión dedicada y un tablero de "
            "indicadores en tiempo real, con el objetivo de alcanzar un 85% de nivel de servicio en 90 días."
        ),
        alcance=[
            "Atención de consultas, reclamos y solicitudes de clientes en canales voz, correo y chat.",
            "Nivel de servicio objetivo: 85% de llamadas atendidas en 20 segundos.",
            "Satisfacción del cliente (CSAT) superior a 4,3 sobre 5.",
            "Cobertura horaria de lunes a domingo, de 08:00 a 22:00.",
        ],
        tecnica=(
            "El servicio opera sobre una plataforma de telefonía en la nube integrada al CRM del banco. Los agentes "
            "reciben capacitación inicial de 3 semanas en productos bancarios y calidad de servicio. Un equipo de "
            "calidad monitorea el 5% de las interacciones y entrega retroalimentación semanal. Se incluye un bot de "
            "primer nivel para consultas de saldo y estado de solicitudes, que reduce el volumen de contactos simples."
        ),
        equipo=[
            ("Agente de atención", 40, "Atención de contactos entrantes"),
            ("Supervisor", 4, "Gestión de turnos y coaching"),
            ("Analista de calidad", 2, "Monitoreo y reportes"),
            ("Gerente de cuenta", 1, "Relación con el cliente y SLAs"),
        ],
        precios=[
            ("Agentes contact center omnicanal", "Agente/mes", 40, "52,000"),
            ("Supervisión y calidad", "Mes", 1, "9,500"),
            ("Plataforma y bot de primer nivel", "Mes", 1, "6,800"),
        ],
        total=68300,
        condiciones=[
            "Contrato de 24 meses con renovación automática.",
            "Facturación mensual a 30 días.",
            "Penalidad del 5% si el nivel de servicio baja de 80% por dos meses consecutivos.",
            "Reajuste anual según IPC.",
        ],
    ),
    dict(
        fname="Propuesta_BPO_Cobranza_Retail_Mercantil.docx",
        title="Propuesta de Gestión de Cobranza Telefónica",
        client="Retail Mercantil Ltda.",
        resumen=(
            "Retail Mercantil mantiene una cartera morosa de 18.000 clientes con una recuperación del 11% en los "
            "últimos 12 meses. NoonDalton propone un servicio de cobranza telefónica y digital segmentado por tramo "
            "de mora, con 25 gestores, campañas por WhatsApp y SMS, y comisiones alineadas a resultados, para "
            "elevar la recuperación al 18% en 6 meses."
        ),
        alcance=[
            "Gestión de cobranza temprana (1–30 días), media (31–90 días) y avanzada (más de 90 días).",
            "Campañas de recordatorio por WhatsApp, SMS y correo.",
            "Negociación de convenios de pago y repactaciones.",
            "Cumplimiento de la normativa de cobranza y protección de datos.",
        ],
        tecnica=(
            "La cartera se segmenta por tramo de mora, monto y comportamiento de pago. Un discador predictivo "
            "maximiza el contacto efectivo y cada gestor trabaja con guiones de negociación validados por el "
            "cliente. Todas las llamadas se graban y se auditan. Se entrega un informe semanal de recuperación por "
            "tramo y un tablero con la tasa de contacto, la promesa de pago y la recuperación efectiva."
        ),
        equipo=[
            ("Gestor de cobranza", 25, "Contacto y negociación"),
            ("Supervisor", 3, "Control de metas y calidad"),
            ("Analista de cartera", 1, "Segmentación y reportes"),
        ],
        precios=[
            ("Gestión de cobranza telefónica", "Gestor/mes", 25, "31,250"),
            ("Campañas digitales WhatsApp/SMS", "Mes", 1, "3,200"),
            ("Comisión por recuperación", "% sobre monto recuperado", "6%", "variable"),
        ],
        total=34450,
        condiciones=[
            "Contrato de 12 meses.",
            "Modelo mixto: tarifa fija más comisión por recuperación efectiva.",
            "Confidencialidad y protección de datos conforme a la normativa vigente.",
            "Informes semanales y comité mensual de resultados.",
        ],
    ),
    dict(
        fname="Propuesta_BPO_Backoffice_Seguros_Pacifico.docx",
        title="Propuesta de Back Office y Procesamiento de Siniestros",
        client="Seguros del Pacífico S.A.",
        resumen=(
            "Seguros del Pacífico procesa manualmente más de 6.000 siniestros al mes, con un tiempo promedio de "
            "liquidación de 12 días y una tasa de error del 7% en la digitación. NoonDalton propone un servicio de "
            "back office con digitación, validación documental y liquidación de siniestros, apoyado en captura "
            "inteligente de documentos, para bajar la liquidación a 5 días y el error bajo 1%."
        ),
        alcance=[
            "Recepción, digitación y validación documental de siniestros.",
            "Liquidación de siniestros de baja complejidad según reglas de negocio.",
            "Conciliación de pagos a proveedores y talleres.",
            "Reportes diarios de volumen, tiempos y errores.",
        ],
        tecnica=(
            "Los documentos ingresan por correo y portal, y una herramienta de reconocimiento óptico (OCR) extrae "
            "los datos que luego valida un analista. Los casos con inconsistencias se escalan a un revisor senior. "
            "El proceso se documenta en manuales operativos y se mide con indicadores de tiempo de ciclo y calidad "
            "de digitación. Se incluye un plan de continuidad operativa con respaldo en un segundo centro."
        ),
        equipo=[
            ("Analista de back office", 18, "Digitación y validación"),
            ("Liquidador", 6, "Liquidación de siniestros"),
            ("Revisor senior", 2, "Control de calidad y excepciones"),
            ("Líder de operación", 1, "Gestión del servicio"),
        ],
        precios=[
            ("Procesamiento de siniestros (back office)", "Analista/mes", 24, "30,000"),
            ("Control de calidad y liderazgo", "Mes", 1, "7,400"),
            ("Captura inteligente de documentos (OCR)", "Mes", 1, "4,500"),
        ],
        total=41900,
        condiciones=[
            "Contrato de 18 meses.",
            "Facturación mensual según volumen procesado, con piso del 80% del volumen acordado.",
            "SLA de liquidación: 90% de siniestros simples en 5 días hábiles.",
            "Auditoría semestral de calidad por parte del cliente.",
        ],
    ),
]


def _table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for cell, h in zip(t.rows[0].cells, headers):
        cell.text = h
    for row in rows:
        for cell, v in zip(t.add_row().cells, row):
            cell.text = str(v)


def build(out: Path, p: dict) -> None:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(p["title"], 0)
    d.add_paragraph(
        f"Preparado para: {p['client']}\nPreparado por: NoonDalton BPO Services\n"
        "Versión 1.0 – Documento confidencial (ejemplo ficticio)"
    )
    d.add_heading("1. Resumen Ejecutivo", 1)
    d.add_paragraph(p["resumen"])
    d.add_heading("2. Descripción de la Empresa", 1)
    d.add_paragraph(EMPRESA)
    d.add_heading("3. Alcance y Objetivos", 1)
    for a in p["alcance"]:
        d.add_paragraph(a, style="List Bullet")
    d.add_heading("4. Descripción del Servicio", 1)
    d.add_paragraph(p["tecnica"])
    d.add_heading("5. Organización del Trabajo", 1)
    _table(d, ["Rol", "Cantidad", "Responsabilidad"], p["equipo"])
    d.add_heading("6. Precios y Condiciones Comerciales", 1)
    _table(d, ["Servicio", "Unidad", "Cantidad", "Precio mensual (USD)"], p["precios"])
    d.add_paragraph(f"Total mensual: USD {p['total']:,}")
    d.add_heading("7. Condiciones Generales", 1)
    for c in p["condiciones"]:
        d.add_paragraph(c, style="List Bullet")
    d.save(out / p["fname"])


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    out.mkdir(parents=True, exist_ok=True)
    for prop in PROPUESTAS:
        build(out, prop)
        print("generado:", out / prop["fname"])
