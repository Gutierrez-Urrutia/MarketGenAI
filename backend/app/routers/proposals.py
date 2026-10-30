import logging
import os
import re
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Optional, Union
from xml.sax.saxutils import escape

import httpx
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import FileResponse, JSONResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.rate_limit import limiter
from app.core.rbac import is_admin
from app.config import settings
from app.dependencies.auth import CurrentUser, get_current_user
from app.schemas.job import JobAccepted
from app.schemas.proposal import (
    GenerateProposalRequest,
    ProposalExportRequest,
    ProposalCreate,
    PROPOSAL_STATUS_VALUES,
    ProposalOut,
    ProposalStatus,
    ProposalUpdate,
    ProposalVersionCreate,
    SendToCrmRequest,
    invalid_status_detail,
)
from app.services.firestore_service import proposals_repo, jobs_repo, settings_repo
from app.services import crm_service, deepseek_service
from app.workers.tasks.proposal_tasks import _lang_code, format_usd
from app.services.crm_service import (
    CrmAuthError,
    CrmError,
    CrmProviderError,
    CrmTimeoutError,
)
from app.rendering import proposal_brand
from app.rendering.proposal_brand import (
    ACCENT as BRAND_ACCENT,
    BORDER as BRAND_BORDER,
    MUTED as BRAND_MUTED,
    PRIMARY as BRAND_PRIMARY,
    PRIMARY_DEEP as BRAND_PRIMARY_DEEP,
    PRIMARY_SOFT as BRAND_PRIMARY_SOFT,
    SURFACE as BRAND_SURFACE,
    TEXT as BRAND_TEXT,
)
from app.workers.tasks.proposal_tasks import (
    _output_dir,
    _format_proposal_line_items,
    _generate_proposal_content,
    task_generate_proposal,
    generate_proposal_now,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/proposals", tags=["Proposals"])

# Brand palette — mirrors the --proposal-* tokens in frontend/src/index.css
# so exported PDFs/DOCX match the on-screen proposal preview.
# ── Helpers internos ──────────────────────────────────────────────────────────


def _assert_owner(proposal: dict, user_id: str):
    """Lanza 403 si la propuesta no pertenece al usuario autenticado."""
    if proposal.get("userId") != user_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Access denied."
        )


async def _get_proposal_for_user(proposal_id: str, user_id: str) -> dict:
    """Obtiene una propuesta de Firestore y verifica ownership. Lanza 404 o 403 según corresponda."""
    try:
        proposal = await proposals_repo.get_or_404(proposal_id)
    except KeyError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Proposal '{proposal_id}' not found."
        )
    _assert_owner(proposal, user_id)
    return proposal


_EXPORT_LABELS = {
    "es": {
        "proposal": "PROPUESTA COMERCIAL", "total": "TOTAL", "issued": "FECHA DE EMISIÓN", "valid": "VÁLIDA HASTA", "prepared": "PREPARADO POR", "reference": "REFERENCIA",
        "total_cost": "Total de la propuesta", "page": "Página", "not_calculated": "No calculado",
        "accept": "Aceptación", "confidential": "Documento confidencial",
        "accept_intro": "Para aceptar esta propuesta, firme y devuelva este documento, o confirme su aceptación por correo electrónico.",
        "sig": ["Nombre", "Cargo", "Fecha", "Firma"], "default_title": "Propuesta Comercial",
    },
    "en": {
        "proposal": "COMMERCIAL PROPOSAL", "total": "TOTAL", "issued": "ISSUED ON", "valid": "VALID UNTIL", "prepared": "PREPARED BY", "reference": "REFERENCE",
        "total_cost": "Proposal total", "page": "Page", "not_calculated": "Not calculated",
        "accept": "Acceptance", "confidential": "Confidential document",
        "accept_intro": "To accept this proposal, sign and return this document, or confirm your acceptance by email.",
        "sig": ["Name", "Title", "Date", "Signature"], "default_title": "Commercial Proposal",
    },
    "pt": {
        "proposal": "PROPOSTA COMERCIAL", "total": "TOTAL", "issued": "DATA DE EMISSÃO", "valid": "VÁLIDA ATÉ", "prepared": "PREPARADO POR", "reference": "REFERÊNCIA",
        "total_cost": "Total da proposta", "page": "Página", "not_calculated": "Não calculado",
        "accept": "Aceitação", "confidential": "Documento confidencial",
        "accept_intro": "Para aceitar esta proposta, assine e devolva este documento, ou confirme a aceitação por e-mail.",
        "sig": ["Nome", "Cargo", "Data", "Assinatura"], "default_title": "Proposta Comercial",
    },
}
_MONTHS = {
    "es": ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"],
    "en": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"],
    "pt": ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"],
}


def _localized_date(value: datetime, lang: str) -> str:
    month = _MONTHS[lang][value.month - 1]
    if lang == "en":
        return f"{month} {value.day}, {value.year}"
    return f"{value.day} de {month} de {value.year}"


def _format_usd(value, not_calculated: str = "Not calculated", lang: str = "en") -> str:
    try:
        amount = float(value)
    except (TypeError, ValueError):
        return not_calculated
    return format_usd(amount, not_calculated, lang)


def _resolve_total_amount(proposal: dict):
    """The frontend has saved this under totalAmount, total_amount, or nested
    in structured(_)Content.totalAmount depending on which flow wrote it —
    check all three rather than assuming one key."""
    structured = proposal.get("structured_content") or proposal.get("structuredContent") or {}
    total_amount = proposal.get("totalAmount")
    if total_amount is None:
        total_amount = proposal.get("total_amount")
    if total_amount is None:
        total_amount = structured.get("totalAmount")
    return total_amount


def _proposal_reference(proposal_id: str | None, issued_on: datetime) -> str:
    """Número de referencia de la portada: NDP-AAAAMMDD-XXXX (fecha de emisión y 4 caracteres del id)."""
    suffix = "".join(ch for ch in str(proposal_id or "") if ch.isalnum())[:4].upper() or "0000"
    return f"NDP-{issued_on:%Y%m%d}-{suffix}"


def _proposal_cover_fields(proposal: dict) -> dict:
    """Pulls the cover-page metadata (client, total, dates...) used by both
    the PDF and DOCX exports, mirroring normalizeProposalDocument() on the frontend."""
    structured = proposal.get("structured_content") or proposal.get("structuredContent") or {}
    total_amount = _resolve_total_amount(proposal)

    created_at = proposal.get("createdAt") or proposal.get("created_at")
    issued_on = None
    if isinstance(created_at, datetime):
        issued_on = created_at
    elif isinstance(created_at, str):
        try:
            issued_on = datetime.fromisoformat(created_at.replace("Z", "+00:00"))
        except ValueError:
            issued_on = None
    issued_on = issued_on or datetime.now(timezone.utc)
    valid_until = issued_on + timedelta(days=30)
    lang = _lang_code(proposal.get("language") or structured.get("language"))
    labels = _EXPORT_LABELS[lang]

    return {
        "client": (
            proposal.get("clientName")
            or proposal.get("customerName")
            or proposal.get("customer_name")
            or structured.get("client")
            or ""
        ),
        "contact": structured.get("contact") or proposal.get("contact") or "",
        # Industria y línea de servicio salían de heurísticas de la plantilla del frontend
        # ("B2B Operations", "NoonDalton Standard Proposal"): no son datos del cliente.
        "industry": "",
        "service_line": "",
        "total_amount_display": _format_usd(total_amount, labels["not_calculated"], lang),
        "issued_on": _localized_date(issued_on, lang),
        "valid_until": _localized_date(valid_until, lang),
        "prepared_by": proposal.get("issuedBy") or "—",
        "reference": _proposal_reference(proposal.get("id"), issued_on),
        "language": lang,
        "labels": labels,
    }


def _flatten_wrapped_tag(soup, tag) -> list:
    """Aplana un bloque del HTML de presentación del frontend (tarjetas, contenedores de tabla)
    a las etiquetas simples que entienden los exportadores."""
    if not getattr(tag, "name", None):
        return []
    if tag.name in ("h2", "p", "ul", "ol", "table"):
        return [tag]
    if tag.name != "div":
        return []
    flat = []
    for child in tag.find_all(["table", "div", "p", "ul", "ol"], recursive=False):
        classes = child.get("class") or []
        if "proposal-detail-card" in classes:
            title = child.find("strong")
            body = child.find("p")
            paragraph = soup.new_tag("p")
            parts = [t.get_text(" ", strip=True) for t in (title, body) if t]
            paragraph.string = ": ".join(part for part in parts if part)
            flat.append(paragraph)
        else:
            flat.extend(_flatten_wrapped_tag(soup, child))
    return flat


def _split_proposal_content(html_content: str):
    """Splits the AI-generated HTML fragment into (content_title, subtitle, body_tags),
    so the leading <h1>/<p> can be promoted into the document cover instead of being
    repeated inside the body. Also accepts the frontend's presentation HTML
    (<article class="proposal-preview-document">), which is what the editor saves."""
    soup = BeautifulSoup(html_content or "", "html.parser")
    paper = soup.find(class_="proposal-paper")
    if paper is not None:
        header = paper.find("header")
        title_tag = header.find("h1") if header else None
        subtitle_tag = header.find(class_="proposal-subtitle") if header else None
        body_tags = []
        for child in paper.find_all(recursive=False):
            classes = child.get("class") or []
            if child.name == "header" or "proposal-total-card" in classes or "proposal-acceptance" in classes:
                continue
            if child.name == "section":
                for sub in child.find_all(recursive=False):
                    body_tags.extend(_flatten_wrapped_tag(soup, sub))
            else:
                body_tags.extend(_flatten_wrapped_tag(soup, child))
        return (
            title_tag.get_text(" ", strip=True) if title_tag else None,
            subtitle_tag.get_text(" ", strip=True) if subtitle_tag else None,
            body_tags,
        )

    body_tags = [
        tag for tag in soup.find_all(["h1", "h2", "p", "ul", "ol", "table"], recursive=False)
    ]

    content_title = None
    if body_tags and body_tags[0].name == "h1":
        content_title = body_tags[0].get_text(" ", strip=True)
        body_tags = body_tags[1:]

    subtitle = None
    if body_tags and body_tags[0].name == "p":
        subtitle = body_tags[0].get_text(" ", strip=True)
        body_tags = body_tags[1:]

    return content_title, subtitle, body_tags


def _is_pricing_table(tag, body_tags) -> bool:
    """El resumen de total va solo bajo la tabla de precios. Los documentos nuevos la marcan con
    class="pricing-table"; en los antiguos (sin marca) la única tabla es la de costos."""
    tables = [t for t in body_tags if t.name == "table"]
    marked = [t for t in tables if "pricing-table" in (t.get("class") or [])]
    return tag in marked if marked else True


def _table_rows_from_html(table_tag):
    header_cells = [th.get_text(" ", strip=True) for th in table_tag.find_all("th")]
    body_rows = []
    for tr in table_tag.find_all("tr"):
        cells = tr.find_all("td")
        if cells:
            body_rows.append([td.get_text(" ", strip=True) for td in cells])
    col_count = len(header_cells) if header_cells else (len(body_rows[0]) if body_rows else 0)
    return header_cells, body_rows, col_count


def _make_pdf_page_decoration(page_label: str, confidential: str):
    def decorate(canvas_obj, doc_template):
        _pdf_page_decoration(canvas_obj, doc_template, page_label, confidential)
    return decorate


def _pdf_page_decoration(canvas_obj, _doc_template, page_label: str = "Page", confidential: str = "Confidential"):
    canvas_obj.saveState()
    canvas_obj.setStrokeColor(colors.HexColor(BRAND_BORDER))
    canvas_obj.setLineWidth(0.5)
    canvas_obj.line(54, 40, LETTER[0] - 54, 40)
    canvas_obj.setFont("Helvetica", 8)
    canvas_obj.setFillColor(colors.HexColor(BRAND_MUTED))
    canvas_obj.drawString(54, 28, f"NoonDalton · {confidential}")
    canvas_obj.drawRightString(LETTER[0] - 54, 28, f"{page_label} {canvas_obj.getPageNumber()}")
    canvas_obj.restoreState()


def _pdf_table(table_tag, styles) -> list:
    header_cells, body_rows, col_count = _table_rows_from_html(table_tag)
    if col_count == 0:
        return []

    data = []
    if header_cells:
        data.append([Paragraph(escape(text), styles["th"]) for text in header_cells])
    for row in body_rows:
        padded = (row + [""] * col_count)[:col_count]
        data.append([Paragraph(escape(text), styles["td"]) for text in padded])

    usable_width = LETTER[0] - 108
    weights = [0.18, 0.36, 0.1, 0.16, 0.2] if col_count == 5 else [1 / col_count] * col_count
    col_widths = [usable_width * weight for weight in weights]

    table = Table(data, colWidths=col_widths, repeatRows=1 if header_cells else 0)
    style_commands = [
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor(BRAND_BORDER)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]
    start_row = 0
    if header_cells:
        style_commands.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_PRIMARY)))
        style_commands.append(("LINEBELOW", (0, 0), (-1, 0), 1.4, colors.HexColor(BRAND_ACCENT)))
        start_row = 1
    for row_index in range(start_row, len(data)):
        if (row_index - start_row) % 2 == 1:
            style_commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor(BRAND_SURFACE)))
    table.setStyle(TableStyle(style_commands))
    return [table]


def _pdf_total_highlight(total_display: str, styles, label: str = "Total Implementation Cost") -> Table:
    usable_width = LETTER[0] - 108
    table = Table(
        [[
            Paragraph(escape(label), styles["total_label"]),
            Paragraph(escape(total_display), styles["total_value"]),
        ]],
        colWidths=[usable_width * 0.6, usable_width * 0.4],
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(BRAND_PRIMARY_SOFT)),
        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor(BRAND_ACCENT)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 14),
        ("RIGHTPADDING", (0, 0), (-1, -1), 14),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    return table


def _pdf_body_flowables(body_tags, styles, total_display: str, total_label: str = "Total Implementation Cost") -> list:
    story = []
    pending_bullets: list[str] | None = None

    def flush_bullets():
        nonlocal pending_bullets
        if pending_bullets:
            story.append(ListFlowable(
                [ListItem(Paragraph(escape(item), styles["body"]), leftIndent=14) for item in pending_bullets],
                bulletType="bullet",
                start="circle",
                leftIndent=16,
                bulletDedent=11,  # la viñeta queda dentro del margen del texto
                bulletFontSize=7,
                bulletColor=colors.HexColor(BRAND_ACCENT),
            ))
            story.append(Spacer(1, 6))
            pending_bullets = None

    for tag in body_tags:
        name = tag.name
        if name == "h2":
            flush_bullets()
            heading_para = Paragraph(escape(tag.get_text(" ", strip=True)), styles["h2"])
            heading_rule = HRFlowable(width="100%", thickness=1.4, color=colors.HexColor(BRAND_ACCENT), spaceAfter=10, hAlign="LEFT")
            heading_para.keepWithNext = True
            heading_rule.keepWithNext = True
            story.append(heading_para)
            story.append(heading_rule)
        elif name == "p":
            flush_bullets()
            text = tag.get_text(" ", strip=True)
            if text:
                story.append(Paragraph(escape(text), styles["body"]))
        elif name == "ul":
            flush_bullets()
            pending_bullets = [li.get_text(" ", strip=True) for li in tag.find_all("li") if li.get_text(strip=True)]
        elif name == "ol":
            flush_bullets()
            items = [li.get_text(" ", strip=True) for li in tag.find_all("li") if li.get_text(strip=True)]
            if items:
                story.append(ListFlowable(
                    [ListItem(Paragraph(escape(item), styles["body"])) for item in items],
                    bulletType="1",
                    leftIndent=16,
                    bulletFontSize=9,
                    bulletColor=colors.HexColor(BRAND_PRIMARY),
                ))
                story.append(Spacer(1, 6))
        elif name == "table":
            flush_bullets()
            table_flowables = _pdf_table(tag, styles)
            if table_flowables:
                story.extend(table_flowables)
                story.append(Spacer(1, 10))
                if _is_pricing_table(tag, body_tags):
                    story.append(_pdf_total_highlight(total_display, styles, total_label))
                    story.append(Spacer(1, 10))
    flush_bullets()
    return story


def _pdf_acceptance_block(styles, labels: dict, usable_width: float) -> list:
    """Bloque de firma: cuadrícula de 2x2 campos, cada uno con su propia línea y la etiqueta debajo.
    Va completo en una misma página (KeepTogether) para que no se parta entre dos."""
    gap = 28
    field_w = (usable_width - 12 - gap) / 2  # 12 = relleno interno del marco de ReportLab
    line_color = colors.HexColor(BRAND_MUTED)
    name, role, date, signature = (Paragraph(escape(text), styles["info_label"]) for text in labels["sig"])
    rows = [
        ["", "", ""],            # espacio para escribir / firmar (fila 0)
        [name, "", role],        # etiquetas (fila 1)
        ["", "", ""],            # separación
        ["", "", ""],            # espacio para escribir / firmar (fila 3)
        [date, "", signature],   # etiquetas (fila 4)
    ]
    table = Table(rows, colWidths=[field_w, gap, field_w], rowHeights=[34, 16, 10, 44, 16], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("LINEABOVE", (0, 1), (0, 1), 0.9, line_color),
        ("LINEABOVE", (2, 1), (2, 1), 0.9, line_color),
        ("LINEABOVE", (0, 4), (0, 4), 0.9, line_color),
        ("LINEABOVE", (2, 4), (2, 4), 0.9, line_color),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    heading = Paragraph(escape(labels["accept"]), styles["h2"])
    rule = HRFlowable(width="100%", thickness=1.4, color=colors.HexColor(BRAND_ACCENT), spaceAfter=10, hAlign="LEFT")
    return [
        Spacer(1, 18),
        KeepTogether([
            heading,
            rule,
            Paragraph(escape(labels["accept_intro"]), styles["body"]),
            Spacer(1, 8),
            table,
        ]),
    ]


def _pdf_logo_band(styles, labels: dict, usable_width: float):
    """Banda de color primario con el logo (blanco) a la izquierda y la etiqueta del documento a la derecha."""
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image

    if proposal_brand.logo_available():
        width, height = ImageReader(str(proposal_brand.LOGO_PATH)).getSize()
        logo_h = 30
        left = Image(str(proposal_brand.LOGO_PATH), width=logo_h * width / height, height=logo_h)
    else:
        left = Paragraph(escape(proposal_brand.NAME), styles["band_name"])
    right = Paragraph(escape(labels["proposal"]), styles["band_label"])
    band = Table([[left, right]], colWidths=[usable_width * 0.5, usable_width * 0.5])
    band.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(BRAND_PRIMARY)),
        ("LINEBELOW", (0, 0), (-1, -1), 3, colors.HexColor(BRAND_ACCENT)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 16),
        ("RIGHTPADDING", (0, 0), (-1, -1), 16),
        ("TOPPADDING", (0, 0), (-1, -1), 14),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
    ]))
    return band


def _build_proposal_pdf(file_path: Path, record_title: str, html_content: str, cover: dict) -> None:
    navy = colors.HexColor(BRAND_PRIMARY)
    navy_deep = colors.HexColor(BRAND_PRIMARY_DEEP)
    accent = colors.HexColor(BRAND_ACCENT)
    muted = colors.HexColor(BRAND_MUTED)
    border = colors.HexColor(BRAND_BORDER)
    text_color = colors.HexColor(BRAND_TEXT)

    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("ProposalTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=navy_deep, alignment=0, spaceAfter=4),
        "subtitle": ParagraphStyle("ProposalSubtitle", parent=base["Normal"], fontSize=11, leading=15, textColor=muted, spaceAfter=12),
        "meta": ParagraphStyle("ProposalMeta", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=navy, spaceAfter=2),
        "h2": ParagraphStyle("ProposalH2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=13.5, leading=16, textColor=navy, spaceBefore=16, spaceAfter=4),
        "body": ParagraphStyle("ProposalBody", parent=base["BodyText"], fontSize=10, leading=15, textColor=text_color, spaceAfter=8),
        "th": ParagraphStyle("ProposalTh", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.white),
        "td": ParagraphStyle("ProposalTd", parent=base["Normal"], fontSize=9, textColor=text_color),
        "info_label": ParagraphStyle("ProposalInfoLabel", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7.5, textColor=muted, spaceAfter=2),
        "info_value": ParagraphStyle("ProposalInfoValue", parent=base["Normal"], fontSize=10, textColor=text_color),
        "band_name": ParagraphStyle("ProposalBandName", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=16, textColor=colors.white),
        "band_label": ParagraphStyle("ProposalBandLabel", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.white, alignment=2),
        "total_label": ParagraphStyle("ProposalTotalLabel", parent=base["Normal"], fontSize=10.5, textColor=navy_deep),
        "total_value": ParagraphStyle("ProposalTotalValue", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=13, textColor=navy_deep, alignment=2),
    }

    content_title, subtitle, body_tags = _split_proposal_content(html_content)
    usable_width = LETTER[0] - 108

    story = [
        _pdf_logo_band(styles, cover["labels"], usable_width),
        Spacer(1, 18),
        Paragraph(escape(content_title or record_title or cover["labels"]["default_title"]), styles["title"]),
    ]
    if subtitle:
        story.append(Paragraph(escape(subtitle), styles["subtitle"]))
    meta_bits = [bit for bit in [cover["industry"], cover["contact"], cover["service_line"]] if bit]
    meta_bits.append(f"{cover['labels']['reference']}: {cover['reference']}")
    story.append(Paragraph(escape("   •   ".join(meta_bits)), styles["meta"]))
    story.append(Spacer(1, 12))

    labels = cover["labels"]
    info_pairs = [
        (labels["total"], cover["total_amount_display"]),
        (labels["issued"], cover["issued_on"]),
        (labels["valid"], cover["valid_until"]),
        (labels["prepared"], cover["prepared_by"]),
    ]
    info_table = Table(
        [[[Paragraph(label, styles["info_label"]), Paragraph(escape(str(value) or "—"), styles["info_value"])] for label, value in info_pairs]],
        colWidths=[usable_width / 4] * 4,
    )
    info_table.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, border),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, border),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 12),
        ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=0.75, color=border, spaceAfter=4))

    story.extend(_pdf_body_flowables(body_tags, styles, cover["total_amount_display"], labels["total_cost"]))
    story.extend(_pdf_acceptance_block(styles, labels, usable_width))

    pdf = SimpleDocTemplate(
        str(file_path),
        pagesize=LETTER,
        leftMargin=54,
        rightMargin=54,
        topMargin=46,
        bottomMargin=56,
    )
    decorate = _make_pdf_page_decoration(labels["page"], labels["confidential"])
    pdf.build(story, onFirstPage=decorate, onLaterPages=decorate)


def _set_run_color(run, hex_color: str) -> None:
    run.font.color.rgb = RGBColor.from_string(hex_color.lstrip("#"))


def _shade_cell(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_borders(cell, hex_color: str, sz: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(sz))
        element.set(qn("w:color"), hex_color.lstrip("#"))
        borders.append(element)
    tc_pr.append(borders)


def _add_heading_rule(paragraph, hex_color: str) -> None:
    """Adds a colored bottom border under a heading paragraph, mirroring the
    h2 underline used in the on-screen proposal preview."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), hex_color.lstrip("#"))
    p_borders.append(bottom)
    p_pr.append(p_borders)


def _add_docx_pricing_table(doc: Document, table_tag) -> None:
    header_cells, body_rows, col_count = _table_rows_from_html(table_tag)
    if col_count == 0:
        return

    table = doc.add_table(rows=0, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if header_cells:
        header_row = table.add_row()
        for col, text in enumerate(header_cells):
            cell = header_row.cells[col]
            cell.text = text
            _shade_cell(cell, BRAND_PRIMARY)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9.5)
                _set_run_color(run, "FFFFFF")

    for row_index, row_values in enumerate(body_rows):
        row = table.add_row()
        padded = (row_values + [""] * col_count)[:col_count]
        for col, text in enumerate(padded):
            cell = row.cells[col]
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9.5)
            if row_index % 2 == 1:
                _shade_cell(cell, BRAND_SURFACE)

    for row in table.rows:
        for cell in row.cells:
            _set_cell_borders(cell, BRAND_BORDER)

    if col_count == 5:  # tabla de precios: servicio, descripción, cantidad, unitario, subtotal
        table.autofit = False
        widths = [Pt(w) for w in (118, 122, 36, 96, 96)]
        for row in table.rows:
            for col, cell in enumerate(row.cells):
                cell.width = widths[col]
                if col >= 2:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()


def _add_docx_total_highlight(doc: Document, total_display: str, label: str = "Total Implementation Cost") -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    label_cell, value_cell = table.rows[0].cells
    label_cell.text = label
    value_cell.text = total_display

    for cell in (label_cell, value_cell):
        _shade_cell(cell, BRAND_PRIMARY_SOFT)
    label_cell._tc.get_or_add_tcPr().append(_left_border_element(BRAND_ACCENT))

    for run in label_cell.paragraphs[0].runs:
        run.font.size = Pt(10.5)
        _set_run_color(run, BRAND_PRIMARY_DEEP)

    value_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in value_cell.paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(13)
        _set_run_color(run, BRAND_PRIMARY_DEEP)

    doc.add_paragraph()


def _left_border_element(hex_color: str):
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:color"), hex_color.lstrip("#"))
    borders = OxmlElement("w:tcBorders")
    borders.append(left)
    return borders


def _set_cell_top_border(cell, hex_color: str, sz: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(sz))
    top.set(qn("w:color"), hex_color.lstrip("#"))
    borders.append(top)
    tc_pr.append(borders)


def _add_docx_signature_grid(doc: Document, names: list) -> None:
    """Cuadrícula de firma 2x2: espacio en blanco, y debajo la etiqueta con una línea superior propia."""
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for left, right in ((names[0], names[1]), (names[2], names[3])):
        space_row = table.add_row()
        space_row.height = Pt(34)
        label_row = table.add_row()
        for row_cells, (l_name, r_name) in ((label_row.cells, (left, right)),):
            for idx, text in ((0, l_name), (2, r_name)):
                cell = row_cells[idx]
                cell.text = text
                _set_cell_top_border(cell, BRAND_MUTED)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(8)
                    _set_run_color(run, BRAND_MUTED)
    for row in table.rows:
        row.cells[0].width = Pt(215)
        row.cells[1].width = Pt(24)
        row.cells[2].width = Pt(215)


def _add_docx_logo_band(doc: Document, labels: dict) -> None:
    """Primera tabla del documento: banda de color primario con el logo y la etiqueta del documento."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    for cell in (left, right):
        _shade_cell(cell, BRAND_PRIMARY)
        cell.width = Pt(234)
    left_par = left.paragraphs[0]
    left_par.paragraph_format.space_before = Pt(6)
    left_par.paragraph_format.space_after = Pt(6)
    if proposal_brand.logo_available():
        left_par.add_run().add_picture(str(proposal_brand.LOGO_PATH), height=Pt(30))
    else:
        run = left_par.add_run(proposal_brand.NAME)
        run.bold = True
        run.font.size = Pt(16)
        _set_run_color(run, "FFFFFF")
    right_par = right.paragraphs[0]
    right_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_par.paragraph_format.space_before = Pt(12)
    run = right_par.add_run(labels["proposal"])
    run.bold = True
    run.font.size = Pt(8.5)
    _set_run_color(run, "FFFFFF")
    # línea de acento bajo la banda
    for cell in (left, right):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "24")
        bottom.set(qn("w:color"), BRAND_ACCENT.lstrip("#"))
        borders.append(bottom)
        tc_pr.append(borders)
    doc.add_paragraph()


def _add_docx_footer(doc: Document, labels: dict) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.text = f"NoonDalton · {labels['confidential']}\t\t{labels['page']} "
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    for item in paragraph.runs:
        item.font.size = Pt(8)
        _set_run_color(item, BRAND_MUTED)


def _build_proposal_docx(file_path: Path, record_title: str, html_content: str, cover: dict) -> None:
    content_title, subtitle, body_tags = _split_proposal_content(html_content)

    doc = Document()
    base_style = doc.styles["Normal"]
    base_style.font.name = "Calibri"
    base_style.font.size = Pt(10.5)
    base_style.font.color.rgb = RGBColor.from_string(BRAND_TEXT.lstrip("#"))

    labels = cover["labels"]
    title_borders = doc.styles["Title"].element.find(".//" + qn("w:pBdr"))
    if title_borders is not None:  # la línea bajo el título trae el azul por defecto de Word
        for edge in title_borders:
            edge.set(qn("w:color"), BRAND_ACCENT.lstrip("#"))
    _add_docx_footer(doc, labels)
    _add_docx_logo_band(doc, labels)
    title_paragraph = doc.add_heading(content_title or record_title or labels["default_title"], level=0)
    for run in title_paragraph.runs:
        run.font.name = "Calibri"
        _set_run_color(run, BRAND_PRIMARY_DEEP)

    if subtitle:
        subtitle_paragraph = doc.add_paragraph(subtitle)
        for run in subtitle_paragraph.runs:
            run.font.size = Pt(10.5)
            _set_run_color(run, BRAND_MUTED)

    meta_bits = [bit for bit in [cover["industry"], cover["contact"], cover["service_line"]] if bit]
    if meta_bits:
        meta_paragraph = doc.add_paragraph("   •   ".join(meta_bits))
        for run in meta_paragraph.runs:
            run.bold = True
            run.font.size = Pt(9)
            _set_run_color(run, BRAND_PRIMARY)

    doc.add_paragraph()

    reference_paragraph = doc.add_paragraph(f"{labels['reference']}: {cover['reference']}")
    reference_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info_table = doc.add_table(rows=2, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_pairs = [
        (labels["total"], cover["total_amount_display"]),
        (labels["issued"], cover["issued_on"]),
        (labels["valid"], cover["valid_until"]),
        (labels["prepared"], cover["prepared_by"]),
    ]
    for col, (label, value) in enumerate(info_pairs):
        label_cell = info_table.cell(0, col)
        label_cell.text = label
        for run in label_cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(7.5)
            _set_run_color(run, BRAND_MUTED)

        value_cell = info_table.cell(1, col)
        value_cell.text = str(value) or "—"
        for run in value_cell.paragraphs[0].runs:
            run.font.size = Pt(10)

        _set_cell_borders(label_cell, BRAND_BORDER)
        _set_cell_borders(value_cell, BRAND_BORDER)

    doc.add_paragraph()

    for tag in body_tags:
        name = tag.name
        if name == "h2":
            heading = doc.add_heading(tag.get_text(" ", strip=True), level=1)
            for run in heading.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(13)
                _set_run_color(run, BRAND_PRIMARY)
            _add_heading_rule(heading, BRAND_ACCENT)
        elif name == "p":
            text = tag.get_text(" ", strip=True)
            if text:
                doc.add_paragraph(text)
        elif name == "ul":
            for li in tag.find_all("li"):
                text = li.get_text(" ", strip=True)
                if text:
                    doc.add_paragraph(text, style="List Bullet")
        elif name == "ol":
            for li in tag.find_all("li"):
                text = li.get_text(" ", strip=True)
                if text:
                    doc.add_paragraph(text, style="List Number")
        elif name == "table":
            _add_docx_pricing_table(doc, tag)
            if _is_pricing_table(tag, body_tags):
                _add_docx_total_highlight(doc, cover["total_amount_display"], labels["total_cost"])

    heading = doc.add_heading(labels["accept"], level=1)
    for run in heading.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        _set_run_color(run, BRAND_PRIMARY)
    _add_heading_rule(heading, BRAND_ACCENT)
    doc.add_paragraph(labels["accept_intro"])
    _add_docx_signature_grid(doc, labels["sig"])

    doc.save(file_path)





# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.get("", response_model=List[ProposalOut])
async def get_proposals(
    search: Optional[str] = Query(None, description="Búsqueda parcial en título o cliente"),
    limit: int = Query(200, ge=1, le=500),
    user: CurrentUser = Depends(get_current_user),
):
    proposals = await proposals_repo.list(
        filters=[("userId", "==", user.sub)],
        order_by="updatedAt",
        order_direction="DESCENDING",
        limit=limit,
    )
    if search:
        s = search.lower()
        proposals = [
            p for p in proposals
            if s in (p.get("title") or "").lower() or s in (p.get("clientName") or "").lower()
        ]
    return proposals


@router.post("", response_model=ProposalOut, status_code=status.HTTP_201_CREATED)
async def create_proposal(
    body: ProposalCreate,
    user: CurrentUser = Depends(get_current_user),
):
    data = body.model_dump(exclude_none=True)

    client_name = (
        data.get("clientName")
        or data.get("client")
        or data.get("company")
        or data.get("customer_name")
        or data.get("customer")
        or "Demo Client"
    )
    description = (
        data.get("description")
        or data.get("brief")
        or data.get("summary")
        or data.get("proposal_description")
        or ""
    )
    content = data.get("content") or description

    proposal = {
        **data,
        "title": data["title"],
        "clientName": client_name,
        "description": description,
        "content": content,
        "downloadUrl": None,
        "filePath": None,
        "status": data.get("status") or "draft",
        # Toda propuesta nace en borrador; se ignora lo que envíe el cliente.
        "proposal_status": ProposalStatus.borrador.value,
        "userId": user.sub,
    }
    if data.get("opportunity_id"):
        proposal["opportunity_id"] = data["opportunity_id"]
    return await proposals_repo.create(proposal, doc_id=str(uuid.uuid4()))


@router.get("/{proposal_id}", response_model=ProposalOut)
async def get_proposal(proposal_id: str, user: CurrentUser = Depends(get_current_user)):
    return await _get_proposal_for_user(proposal_id, user.sub)


@router.post(
    "/generate-draft",
    response_model=Union[JobAccepted, ProposalOut],
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit("20/minute")
async def generate_proposal_draft(
    request: Request,
    body: GenerateProposalRequest,
    user: CurrentUser = Depends(get_current_user),
):
    extra = body.model_extra or {}
    title = body.title or extra.get("name") or "AI Marketing Proposal"
    client_name = (
        body.clientName
        or extra.get("client")
        or extra.get("company")
        or extra.get("customer_name")
        or extra.get("customer")
        or "Demo Client"
    )
    description = (
        body.description
        or extra.get("brief")
        or extra.get("summary")
        or extra.get("proposalDescription")
        or extra.get("proposal_description")
        or ""
    )

    proposal = await proposals_repo.create({
        "title": title,
        "clientName": client_name,
        "description": description,
        "content": "",
        "downloadUrl": None,
        "filePath": None,
        "status": "pending",
        "proposal_status": ProposalStatus.borrador.value,
        "language": body.language,
        "customPrompt": body.customPrompt,
        "userId": user.sub,
    }, doc_id=str(uuid.uuid4()))

    options = body.model_dump(exclude_none=True)
    if extra:
        options.update(extra)
    options["userId"] = user.sub  # tras extra: el cliente no puede suplantarlo; el worker lo usa para recuperar el contexto RAG del usuario

    job = await jobs_repo.create_job(
        "proposal_generation",
        user.sub,
        {"proposalId": proposal["id"], "title": title},
    )

    if getattr(body, "sync", False):
        try:
            return await generate_proposal_now(
                job["id"], proposal["id"], title, client_name, description, options
            )
        except Exception as exc:
            await proposals_repo.update(proposal["id"], {"status": "error"})
            await jobs_repo.fail_job(job["id"], str(exc))
            raise

    try:
        task_generate_proposal.delay(
            job["id"], proposal["id"], title, client_name, description, options
        )
        return JSONResponse(
            status_code=status.HTTP_202_ACCEPTED,
            content={
                "job_id": job["id"],
                "proposal_id": proposal["id"],
                "status": "pending",
                "message": "Job enqueued successfully",
            },
        )
    except Exception as exc:
        logger.warning("Celery dispatch failed: %s; falling back to inline execution for job %s", exc, job["id"])
        try:
            return await generate_proposal_now(
                job["id"], proposal["id"], title, client_name, description, options
            )
        except Exception as exc2:
            await proposals_repo.update(proposal["id"], {"status": "error"})
            await jobs_repo.fail_job(job["id"], str(exc2))
            raise


@router.post(
    "/{proposal_id}/generate",
    response_model=Union[JobAccepted, ProposalOut],
    status_code=status.HTTP_200_OK,
)
@limiter.limit("20/minute")
async def generate_proposal_by_id(
    request: Request,
    proposal_id: str,
    body: GenerateProposalRequest,
    user: CurrentUser = Depends(get_current_user),
):
    proposal = await _get_proposal_for_user(proposal_id, user.sub)

    title = body.title or proposal.get("title") or "AI Marketing Proposal"
    client_name = body.clientName or proposal.get("clientName") or "Demo Client"
    description = body.description or proposal.get("description") or ""

    options = body.model_dump(exclude_none=True)
    if body.model_extra:
        options.update(body.model_extra)
    options["userId"] = user.sub  # tras extra: el cliente no puede suplantarlo; el worker lo usa para recuperar el contexto RAG del usuario

    job = await jobs_repo.create_job(
        "proposal_generation",
        user.sub,
        {"proposalId": proposal_id, "title": title},
    )

    if getattr(body, "sync", False):
        try:
            return await generate_proposal_now(
                job["id"], proposal_id, title, client_name, description, options
            )
        except Exception as exc:
            await proposals_repo.update(proposal_id, {"status": "error"})
            await jobs_repo.fail_job(job["id"], str(exc))
            raise

    try:
        task_generate_proposal.delay(
            job["id"], proposal_id, title, client_name, description, options
        )
        return JSONResponse(
            status_code=status.HTTP_202_ACCEPTED,
            content={
                "job_id": job["id"],
                "proposal_id": proposal_id,
                "status": "pending",
                "message": "Job enqueued successfully",
            },
        )
    except Exception as exc:
        logger.warning("Celery dispatch failed: %s; falling back to inline execution for job %s", exc, job["id"])
        try:
            return await generate_proposal_now(
                job["id"], proposal_id, title, client_name, description, options
            )
        except Exception as exc2:
            await proposals_repo.update(proposal_id, {"status": "error"})
            await jobs_repo.fail_job(job["id"], str(exc2))
            raise


@router.put("/{proposal_id}", response_model=ProposalOut)
async def update_proposal(
    proposal_id: str,
    body: ProposalUpdate,
    user: CurrentUser = Depends(get_current_user),
):
    proposal = await _get_proposal_for_user(proposal_id, user.sub)
    update_data = body.model_dump(exclude_none=True)
    if not update_data:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No fields to update."
        )
    # El estado comercial solo cambia por PATCH /status (transiciones + auditoría).
    # `status` (ciclo de generación con IA) sí pasa: el frontend lo usa.
    update_data.pop("proposal_status", None)
    if not update_data:
        return proposal
    return await proposals_repo.update(proposal_id, update_data)


@router.patch("/{proposal_id}/status")
async def update_proposal_status(
    proposal_id: str,
    body: dict,
    request: Request,
    current_user: CurrentUser = Depends(get_current_user),
):
    new_status = body.get("status")
    if new_status not in PROPOSAL_STATUS_VALUES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=invalid_status_detail(new_status),
        )

    actor_is_admin = is_admin(current_user)
    if not actor_is_admin:
        # Igual que antes para los no admin: 404 / 403 si la propuesta no es suya.
        # Un admin puede llegar a una propuesta ajena, pero solo para reabrirla:
        # lo decide la transacción (decide_cambio_estado).
        await _get_proposal_for_user(proposal_id, current_user.sub)

    outcome = await proposals_repo.change_status(
        proposal_id,
        new_status,
        actor_id=current_user.sub,
        actor_roles=current_user.roles,
        actor_is_admin=actor_is_admin,
        method=request.method,
        path=request.url.path,
        ip=request.client.host if request.client else None,
    )
    if outcome.get("not_found"):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Proposal '{proposal_id}' not found.",
        )
    if outcome.get("error"):
        error = outcome["error"]
        raise HTTPException(status_code=error.status_code, detail=error.detail)
    return outcome["proposal"]


@router.delete("/{proposal_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_proposal(proposal_id: str, user: CurrentUser = Depends(get_current_user)):
    await _get_proposal_for_user(proposal_id, user.sub)
    await proposals_repo.delete(proposal_id)
    return None


@router.get("/{proposal_id}/versions")
async def list_proposal_versions(proposal_id: str, user: CurrentUser = Depends(get_current_user)):
    proposal = await _get_proposal_for_user(proposal_id, user.sub)
    versions = proposal.get("versions") or []
    if not versions:
        versions = [{
            "version": 1,
            "title": proposal.get("title"),
            "content": proposal.get("content") or proposal.get("description") or "",
            "createdAt": proposal.get("createdAt"),
            "label": "Current",
        }]
    return {"items": versions, "total": len(versions)}


@router.post("/{proposal_id}/versions")
async def create_proposal_version(
    proposal_id: str,
    body: ProposalVersionCreate,
    user: CurrentUser = Depends(get_current_user),
):
    proposal = await _get_proposal_for_user(proposal_id, user.sub)
    versions = list(proposal.get("versions") or [])
    version = {
        "version": len(versions) + 1,
        "title": body.title or proposal.get("title"),
        "content": body.content or proposal.get("content") or "",
        "label": body.label or f"Version {len(versions) + 1}",
        "createdAt": datetime.now(timezone.utc).isoformat(),
    }
    versions.append(version)
    await proposals_repo.update(proposal_id, {"versions": versions})
    return version


_CRM_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


async def _record_crm_failure(proposal_id: str, existing_sync: dict, provider: str, error: str, partial: dict) -> None:
    """Registra el fallo en crmSync conservando los IDs del CRM ya vinculados (y los creados
    a medias en este intento), para que un reenvío actualice en vez de duplicar."""
    await proposals_repo.update(proposal_id, {"crmSync": {
        **existing_sync,
        **partial,
        "provider": provider,
        "status": "failed",
        "error": error,
        "syncedAt": datetime.now(timezone.utc).isoformat(),
    }})


@router.post("/{proposal_id}/send-to-crm")
async def send_proposal_to_crm(
    proposal_id: str,
    body: SendToCrmRequest,
    user: CurrentUser = Depends(get_current_user),
):
    proposal = await _get_proposal_for_user(proposal_id, user.sub)

    client_email = (proposal.get("clientEmail") or "").strip()
    if not client_email:
        raise HTTPException(
            status_code=400,
            detail="La propuesta no tiene un email de cliente. Edítala para agregarlo antes de enviarla al CRM.",
        )
    if not _CRM_EMAIL_RE.match(client_email):
        raise HTTPException(
            status_code=400,
            detail="El email del cliente no tiene un formato válido. Corrígelo antes de enviarla al CRM.",
        )

    requested_provider = (body.provider or "").strip().lower() or None
    if requested_provider and requested_provider not in ("hubspot", "salesforce"):
        raise HTTPException(status_code=400, detail=f"Proveedor '{requested_provider}' no soportado aún")

    org_settings = await settings_repo.get_by_user(user.sub)
    crm_settings = org_settings.get("crm") or {}
    connections = crm_settings.get("connections") or {}
    existing_sync = proposal.get("crmSync") or {}
    total_amount = _resolve_total_amount(proposal)

    # The frontend doesn't always know which provider is configured (e.g. the
    # Proposals view has no access to the Settings page's React state), so
    # fall back to whichever provider the user marked active in
    # Settings -> Integrations.
    provider = requested_provider or crm_settings.get("activeProvider") or "none"
    if provider == "none":
        raise HTTPException(status_code=400, detail="No hay una integración de CRM configurada en Ajustes.")
    if provider not in ("hubspot", "salesforce"):
        raise HTTPException(status_code=400, detail=f"Proveedor '{provider}' no soportado aún")

    client = crm_service.build_client_from_stored(provider, connections)
    if client is None:
        raise HTTPException(status_code=400, detail="No hay una integración de CRM configurada en Ajustes.")

    partial: dict = {}
    try:
        if provider == "hubspot":
            hubspot = client
            contact_id = await hubspot.upsert_contact(
                email=client_email,
                name=proposal.get("clientName"),
                company=proposal.get("clientCompany"),
            )
            partial["hubspotContactId"] = contact_id
            deal_id = await hubspot.upsert_deal(
                deal_id=existing_sync.get("hubspotDealId"),
                dealname=proposal.get("title"),
                amount=total_amount,
                contact_id=contact_id,
            )
            crm_sync = {
                "provider": "hubspot",
                "status": "synced",
                "syncedAt": datetime.now(timezone.utc).isoformat(),
                "hubspotContactId": contact_id,
                "hubspotDealId": deal_id,
            }
        else:
            salesforce = client
            account_name = proposal.get("clientCompany") or proposal.get("clientName") or "Untitled Account"
            account_id = await salesforce.upsert_account(account_name)
            partial["salesforceAccountId"] = account_id
            contact_id = await salesforce.upsert_contact(
                email=client_email,
                name=proposal.get("clientName"),
                account_id=account_id,
            )
            partial["salesforceContactId"] = contact_id
            opportunity_id = await salesforce.upsert_opportunity(
                opportunity_id=existing_sync.get("salesforceOpportunityId"),
                name=proposal.get("title"),
                amount=total_amount,
                account_id=account_id,
                contact_id=contact_id,
            )
            crm_sync = {
                "provider": "salesforce",
                "status": "synced",
                "syncedAt": datetime.now(timezone.utc).isoformat(),
                "salesforceAccountId": account_id,
                "salesforceContactId": contact_id,
                "salesforceOpportunityId": opportunity_id,
            }
    except CrmError as exc:
        partial.update(getattr(exc, "partial_ids", None) or {})
        if isinstance(exc, CrmAuthError):
            code = 401
        elif isinstance(exc, CrmTimeoutError):
            code = 504
        else:
            code = 502
        await _record_crm_failure(proposal_id, existing_sync, provider, str(exc), partial)
        raise HTTPException(status_code=code, detail=str(exc)) from exc
    except httpx.HTTPError as exc:
        # CRM caído / sin red / DNS: antes escapaba como 500 sin mensaje.
        label = "HubSpot" if provider == "hubspot" else "Salesforce"
        detail = f"No se pudo conectar con {label}. Revisa que el servicio esté disponible e inténtalo de nuevo."
        logger.warning("send-to-crm: fallo de conexión con %s: %r", label, exc)
        await _record_crm_failure(proposal_id, existing_sync, provider, detail, partial)
        raise HTTPException(status_code=502, detail=detail) from exc

    return await proposals_repo.update(proposal_id, {"crmSync": crm_sync})


@router.post("/{proposal_id}/export")
async def export_proposal(
    proposal_id: str,
    body: ProposalExportRequest,
    user: CurrentUser = Depends(get_current_user),
):
    """Exporta la propuesta a PDF/DOCX con texto real. Acepta el contenido editado en pantalla
    (no se guarda: solo se usa para este archivo)."""
    proposal = await _get_proposal_for_user(proposal_id, user.sub)
    html_content = body.content if body.content else proposal.get("content", "")
    executive = body.issuedBy or user.name or user.preferred_username or user.email
    return _render_proposal_file(proposal, proposal_id, body.format, html_content, executive)


@router.get("/{proposal_id}/download")
async def download_proposal(
    proposal_id: str,
    format: str = Query("pdf", pattern="^(pdf|docx)$"),
    user: CurrentUser = Depends(get_current_user),
):
    proposal = await _get_proposal_for_user(proposal_id, user.sub)
    executive = user.name or user.preferred_username or user.email
    return _render_proposal_file(proposal, proposal_id, format, proposal.get("content", ""), executive)


def _render_proposal_file(proposal: dict, proposal_id: str, format: str, html_content: str, issued_by: str | None = None):
    title = proposal.get("title") or "proposal"
    cover = _proposal_cover_fields({**proposal, "id": proposal.get("id") or proposal_id, "issuedBy": issued_by or proposal.get("issuedBy")})

    output_dir = _output_dir()
    safe_title = title.replace(" ", "_").replace("/", "_").lower()

    if format == "docx":
        file_path = output_dir / f"{safe_title}-{proposal_id}.docx"
        _build_proposal_docx(file_path, title, html_content, cover)

        return FileResponse(
            path=file_path,
            filename=file_path.name,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    file_path = output_dir / f"{safe_title}-{proposal_id}.pdf"
    _build_proposal_pdf(file_path, title, html_content, cover)

    return FileResponse(
        path=file_path,
        filename=file_path.name,
        media_type="application/pdf",
    )
