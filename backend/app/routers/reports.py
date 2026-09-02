"""Reports router — aggregated KPIs and export."""
from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.dependencies.auth import CurrentUser, get_current_user
from app.rendering.brand_styles import (
    BRAND_BORDER,
    BRAND_DEEP_NAVY,
    BRAND_MUTED,
    BRAND_PRIMARY_INDIGO,
    BRAND_SOFT_INDIGO_SURFACE,
    BRAND_TEXT,
)
from app.services.firestore_service import (
    books_repo, proposals_repo, customers_repo, opportunities_repo,
)

router = APIRouter(prefix="/reports", tags=["Reports"])


def _proposal_amount(proposal: dict) -> float:
    amount = (
        proposal.get("totalAmount") or
        proposal.get("total") or
        proposal.get("total_amount") or
        (proposal.get("structured_content") or {}).get("totalAmount") or
        0
    )
    return float(amount or 0)


def _build_books_report_pdf(rows: list[dict]) -> bytes:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=27,
        textColor=colors.HexColor(BRAND_DEEP_NAVY),
        alignment=0,
        spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor(BRAND_MUTED),
        spaceAfter=16,
    )
    cell_style = ParagraphStyle(
        "ReportCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor(BRAND_TEXT),
    )
    header_style = ParagraphStyle(
        "ReportHeader",
        parent=cell_style,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )
    headers = [
        ("Titulo", ("Titulo", "Título", "TÃ­tulo")),
        ("Estado", ("Estado",)),
        ("Tipo", ("Tipo",)),
        ("Capitulos", ("Capitulos", "Capítulos", "CapÃ­tulos")),
        ("Actualizado", ("Actualizado",)),
    ]
    data = [[Paragraph(label, header_style) for label, _keys in headers]]
    for row in rows:
        data.append([
            Paragraph(str(next((row.get(key) for key in keys if row.get(key) is not None), "")), cell_style)
            for _label, keys in headers
        ])
    if len(data) == 1:
        data.append([Paragraph("Sin libros para reportar.", cell_style), "", "", "", ""])

    table = Table(data, colWidths=[170, 75, 95, 55, 110], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_PRIMARY_INDIGO)),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(BRAND_SOFT_INDIGO_SURFACE)]),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(BRAND_BORDER)),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor(BRAND_BORDER)),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=42, rightMargin=42, topMargin=44, bottomMargin=44)
    story = [
        Paragraph("Books Report", title_style),
        Paragraph(f"NoonDalton AI Marketing Suite - {datetime.now(timezone.utc).date().isoformat()}", meta_style),
        table,
        Spacer(1, 8),
        Paragraph("Generated from the authenticated user's book library.", meta_style),
    ]
    doc.build(story)
    return buf.getvalue()


@router.get("/overview")
async def reports_overview(
    days: Optional[int] = None,
    user: CurrentUser = Depends(get_current_user),
):
    """Return high-level KPIs for the current user."""
    # Run counts in parallel (Firestore async)
    import asyncio
    total_books, total_proposals, total_customers = await asyncio.gather(
        books_repo.count(filters=[("userId", "==", user.sub)]),
        proposals_repo.count(filters=[("userId", "==", user.sub)]),
        customers_repo.count(filters=[("userId", "==", user.sub)]),
    )

    # Books — chapters count via listing
    books = await books_repo.list(
        filters=[("userId", "==", user.sub)], limit=200,
    )
    chapters_generated = sum(b.get("chapterCount") or 0 for b in books)
    proposals = await proposals_repo.list(filters=[("userId", "==", user.sub)], limit=500)
    total_proposal_value = sum(_proposal_amount(proposal) for proposal in proposals)
    opportunities = await opportunities_repo.list(filters=[("userId", "==", user.sub)], limit=500)
    cutoff = datetime.now(timezone.utc) - timedelta(days=days) if days else None
    if cutoff:
        def within_period(opportunity: dict) -> bool:
            created_at = opportunity.get("createdAt")
            if not created_at:
                return False
            if isinstance(created_at, datetime):
                value = created_at
            else:
                try:
                    value = datetime.fromisoformat(str(created_at).replace("Z", "+00:00"))
                except ValueError:
                    return False
            if value.tzinfo is None:
                value = value.replace(tzinfo=timezone.utc)
            return value > cutoff

        opportunities = [o for o in opportunities if within_period(o)]
    won_stages = {"Won", "Customer"}
    qualified_stages = {"Contacted", "Replied", "In Conversation", "Won", "Customer", "Lost"}
    won_count = sum(1 for o in opportunities if o.get("stage") in won_stages)
    qualified_count = sum(1 for o in opportunities if o.get("stage") in qualified_stages)
    win_rate = round((won_count / qualified_count * 100), 1) if qualified_count > 0 else 0

    # Status breakdown
    status_map: dict[str, int] = {}
    for b in books:
        status_map[b.get("status", "draft")] = status_map.get(b.get("status", "draft"), 0) + 1

    books_by_status = [{"name": k, "value": v} for k, v in status_map.items()]

    return {
        "kpis": {
            "totalBooks":       total_books,
            "chaptersGenerated": chapters_generated,
            "proposalsSent":    total_proposals,
            "totalProposalValue": total_proposal_value,
            "winRate": win_rate,
            "totalCustomers":   total_customers,
        },
        "win_rate": win_rate,
        "win_rate_detail": f"{won_count} won / {qualified_count} qualified",
        "booksByStatus": books_by_status,
        "generatedAt":   datetime.now(timezone.utc).isoformat(),
    }

@router.get("/dashboard")
async def reports_dashboard(user: CurrentUser = Depends(get_current_user)):
    """Dashboard KPIs for the web app, scoped to the authenticated user."""
    import asyncio

    total_books, total_proposals, total_customers = await asyncio.gather(
        books_repo.count(filters=[("userId", "==", user.sub)]),
        proposals_repo.count(filters=[("userId", "==", user.sub)]),
        customers_repo.count(filters=[("userId", "==", user.sub)]),
    )
    proposals = await proposals_repo.list(filters=[("userId", "==", user.sub)], limit=500)
    accepted = len([p for p in proposals if p.get("status") in {"accepted", "won"}])
    sent = len([p for p in proposals if p.get("status") in {"sent", "generated", "accepted", "won"}])
    return {
        "detected": total_customers + total_proposals + total_books,
        "researched": total_customers,
        "contacted": sent,
        "pendingReview": len([p for p in proposals if p.get("status", "draft") == "draft"]),
        "replied": len([p for p in proposals if p.get("status") == "generated"]),
        "won": accepted,
        "totals": {
            "books": total_books,
            "proposals": total_proposals,
            "customers": total_customers,
            "winRate": round((accepted / sent) * 100, 1) if sent else 0,
        },
        "generatedAt": datetime.now(timezone.utc).isoformat(),
    }

@router.get("/books")
async def reports_books(
    limit:  int = Query(50, ge=1, le=200),
    user:   CurrentUser = Depends(get_current_user),
):
    """Return per-book stats."""
    books = await books_repo.list(
        filters=[("userId", "==", user.sub)],
        order_by="updatedAt", order_direction="DESCENDING",
        limit=limit,
    )
    return {"items": books, "total": len(books)}

@router.get("/proposals")
async def reports_proposals(
    limit: int = Query(10, ge=1, le=100),
    user: CurrentUser = Depends(get_current_user),
):
    proposals = await proposals_repo.list(
        filters=[("userId", "==", user.sub)],
        order_by="createdAt",
        order_direction="DESCENDING",
        limit=500,
    )

    status_map = {}

    for proposal in proposals:
        status = proposal.get("proposal_status") or "Generada"
        amount = _proposal_amount(proposal)

        if status not in status_map:
            status_map[status] = {"label": status, "count": 0, "totalValue": 0}

        status_map[status]["count"] += 1
        status_map[status]["totalValue"] += amount

    return {
        "dimension": "proposal_status",
        "data": list(status_map.values())[:limit],
        "timeSeries": [],
    }

@router.get("/content")
async def reports_content(
    limit: int = Query(50, ge=1, le=200),
    user: CurrentUser = Depends(get_current_user),
):
    books = await books_repo.list(
        filters=[("userId", "==", user.sub)],
        order_by="updatedAt",
        order_direction="DESCENDING",
        limit=limit,
    )

    return {
        "data": [
            {
                "bookId": book.get("id"),
                "bookTitle": book.get("title", ""),
                "status": book.get("status", "draft"),
                "assetCount": book.get("assetCount", 0),
                "downloadCount": book.get("downloadCount", 0),
                "createdAt": str(book.get("createdAt", "")),
            }
            for book in books
        ]
    }

@router.get("/export")
async def export_report(
    format: str = Query("xlsx", pattern="^(xlsx|csv|pdf)$"),
    user:   CurrentUser = Depends(get_current_user),
):
    """Export a simple overview report as Excel, CSV, or PDF."""
    books = await books_repo.list(
        filters=[("userId", "==", user.sub)], limit=500,
    )

    rows = [
        {
            "Título":     b.get("title", ""),
            "Estado":     b.get("status", ""),
            "Tipo":       b.get("contentType", ""),
            "Capítulos":  b.get("chapterCount", 0),
            "Actualizado": str(b.get("updatedAt", "")),
        }
        for b in books
    ]

    if format == "csv":
        import csv
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)
        return StreamingResponse(
            io.BytesIO(buf.getvalue().encode()),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=report.csv"},
        )

    if format == "pdf":
        return StreamingResponse(
            io.BytesIO(_build_books_report_pdf(rows)),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=report.pdf"},
        )

    # xlsx via openpyxl
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return {"error": "openpyxl not installed"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Libros"

    # Header
    headers = ["Título", "Estado", "Tipo", "Capítulos", "Actualizado"]
    header_fill = PatternFill("solid", fgColor="6366F1")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = 25

    for row_idx, row in enumerate(rows, 2):
        for col_idx, key in enumerate(headers, 1):
            ws.cell(row=row_idx, column=col_idx, value=row[key])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=report.xlsx"},
    )
