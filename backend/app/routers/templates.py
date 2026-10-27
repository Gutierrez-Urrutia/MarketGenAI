"""Templates router — manage reusable prompt/content templates.

Plantillas de propuesta (RF-VTA-07): se clasifican por línea de servicio e industria,
traen un bloque de condiciones comerciales y, al aplicarse, completan las variables
estándar con datos de la propuesta y del cliente."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Optional
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from docx import Document

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - dependency is declared in requirements.txt
    PdfReader = None

from app.dependencies.auth import CurrentUser, get_current_user
from app.routers.customers import _get_customer_for_user
from app.routers.proposals import (
    _get_proposal_for_user, _localized_date, _proposal_cover_fields, _resolve_total_amount,
)
from app.schemas.template import (
    TemplateCreate, TemplateUpdate, TemplateOut, TemplateListResponse,
    ApplyTemplateRequest, ApplyTemplateResponse, Industry, ServiceType, TemplateType,
)
from app.services.firestore_service import templates_repo
from app.services.template_service import (
    catalog, classification_labels, default_commercial_terms, lang_code,
    render_placeholders, with_defaults,
)

router = APIRouter(prefix="/templates", tags=["Templates"])

SUPPORTED_UPLOAD_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".docx", ".pdf"}


def _assert_owner(template: dict, user: CurrentUser) -> None:
    if template.get("userId") != user.sub and not template.get("isPublic"):
        raise HTTPException(status_code=403, detail="Access denied.")


async def _get_template_for_user(template_id: str, user: CurrentUser) -> dict:
    template = await templates_repo.get(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Template not found.")
    _assert_owner(template, user)
    return with_defaults(template)


def _matches(template: dict, service_type: Optional[ServiceType], industry: Optional[Industry], search: Optional[str]) -> bool:
    if service_type and template["serviceType"] != service_type.value:
        return False
    if industry and template["industry"] != industry.value:
        return False
    if search:
        s = search.lower()
        return s in template.get("name", "").lower() or s in (template.get("description") or "").lower()
    return True


async def _variables_from_sources(body: ApplyTemplateRequest, template: dict, user: CurrentUser) -> dict[str, Any]:
    """Variables estándar desde la propuesta y el cliente. El cliente del CRM pisa la
    copia que guarda la propuesta, y los valores manuales de `body.variables` pisan todo."""
    values: dict[str, Any] = {}
    language = body.language
    service_type = template["serviceType"]
    industry = template["industry"]

    if body.proposalId:
        proposal = await _get_proposal_for_user(body.proposalId, user.sub)
        cover = _proposal_cover_fields(proposal)
        language = language or cover["language"]
        service_type = proposal.get("serviceType") or service_type
        industry = proposal.get("industry") or industry
        values.update({
            "propuesta_titulo": proposal.get("title"),
            "cliente_nombre":   cover["client"],
            "cliente_empresa":  proposal.get("clientCompany"),
            "cliente_email":    proposal.get("clientEmail"),
            "monto_total":      cover["total_amount_display"] if _resolve_total_amount(proposal) is not None else None,
            "fecha":            cover["issued_on"],
            "valida_hasta":     cover["valid_until"],
            "referencia":       cover["reference"],
            "preparado_por":    proposal.get("issuedBy"),
        })

    if body.customerId:
        customer = await _get_customer_for_user(body.customerId, user)
        values.update({
            "cliente_nombre":  customer.get("name"),
            "cliente_empresa": customer.get("company"),
            "cliente_email":   customer.get("email"),
        })

    lang = lang_code(language)
    values.update(classification_labels(service_type, industry, lang))
    values["fecha"] = values.get("fecha") or _localized_date(datetime.now(timezone.utc), lang)
    values.update(body.variables)
    return values


# ── Catalog ───────────────────────────────────────────────────────────────────
# Declarado antes de /{template_id} para que "catalog" no se tome como id.
@router.get("/catalog")
async def get_template_catalog(user: CurrentUser = Depends(get_current_user)):
    """Líneas de servicio, industrias, variables estándar y condiciones comerciales base."""
    return catalog()


def _decode_text_file(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n\n".join(parts)


def _extract_pdf_text(content: bytes) -> str:
    if PdfReader is None:
        raise HTTPException(status_code=500, detail="PDF extraction is not available. Install pypdf to enable PDF uploads.")
    reader = PdfReader(BytesIO(content))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _extract_template_content(filename: str, content: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension in {".txt", ".md", ".html", ".htm"}:
        return _decode_text_file(content)
    if extension == ".docx":
        return _extract_docx_text(content)
    if extension == ".pdf":
        return _extract_pdf_text(content)
    supported = ", ".join(sorted(SUPPORTED_UPLOAD_EXTENSIONS))
    raise HTTPException(status_code=400, detail=f"Unsupported template file type. Supported formats: {supported}.")


def _template_name_from_filename(filename: str) -> str:
    stem = Path(filename or "template").stem.strip()
    return stem or "template"


@router.post("/extract")
async def extract_template_file(
    file: UploadFile = File(...),
    user: CurrentUser = Depends(get_current_user),
):
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded template file is empty.")
    extracted = _extract_template_content(file.filename or "", content).strip()
    if not extracted:
        raise HTTPException(status_code=400, detail="Could not extract text from this template file.")
    return {
        "name": _template_name_from_filename(file.filename or "template"),
        "content": extracted,
    }


# ── List ──────────────────────────────────────────────────────────────────────
@router.get("", response_model=TemplateListResponse)
async def list_templates(
    template_type: Optional[str] = Query(None, alias="type"),
    service_type:  Optional[ServiceType] = Query(None, alias="serviceType"),
    industry:      Optional[Industry]    = Query(None),
    search:        Optional[str] = Query(None),
    include_public: bool         = Query(True),
    limit:         int           = Query(20, ge=1, le=100),
    offset:        int           = Query(0, ge=0),
    user:          CurrentUser   = Depends(get_current_user),
):
    filters = [("userId", "==", user.sub)]
    if template_type:
        filters.append(("type", "==", template_type))

    # Servicio e industria se filtran en memoria: las plantillas anteriores a RF-VTA-07
    # no tienen esos campos en Firestore y solo los reciben en with_defaults. Con
    # filtros + order_by el repo ya trae la lista completa y ordena en memoria, así
    # que pedirla sin límite no cambia la consulta.
    user_templates = [
        t for t in map(with_defaults, await templates_repo.list(
            filters=filters, order_by="usageCount", order_direction="DESCENDING", limit=0,
        ))
        if _matches(t, service_type, industry, search)
    ]
    total = len(user_templates)
    user_templates = user_templates[offset:offset + limit]

    public_templates = []
    if include_public:
        pub_filters = [("isPublic", "==", True)]
        if template_type:
            pub_filters.append(("type", "==", template_type))
        public_templates = await templates_repo.list(
            filters=pub_filters, order_by="usageCount", order_direction="DESCENDING", limit=0,
        )
        # Deduplicate: remove public templates the user owns (they are already listed or filtered out)
        public_templates = [
            t for t in map(with_defaults, public_templates)
            if t.get("userId") != user.sub and _matches(t, service_type, industry, search)
        ][:10]

    return TemplateListResponse(items=user_templates + public_templates, total=total)


# ── Create ────────────────────────────────────────────────────────────────────
@router.post("", response_model=TemplateOut, status_code=201)
async def create_template(
    body: TemplateCreate,
    user: CurrentUser = Depends(get_current_user),
):
    data = body.model_dump(mode="json")
    if body.type == TemplateType.proposal and not (body.commercialTerms or "").strip():
        data["commercialTerms"] = default_commercial_terms()
    data["userId"] = user.sub
    data["usageCount"] = 0
    return with_defaults(await templates_repo.create(data))


# ── Get one ───────────────────────────────────────────────────────────────────
@router.get("/{template_id}", response_model=TemplateOut)
async def get_template(
    template_id: str,
    user:        CurrentUser = Depends(get_current_user),
):
    return await _get_template_for_user(template_id, user)


# ── Update ────────────────────────────────────────────────────────────────────
@router.put("/{template_id}", response_model=TemplateOut)
async def update_template(
    template_id: str,
    body:        TemplateUpdate,
    user:        CurrentUser = Depends(get_current_user),
):
    template = await templates_repo.get(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Template not found.")
    if template.get("userId") != user.sub:
        raise HTTPException(status_code=403, detail="Only the owner can edit this template.")
    if (
        template.get("type") == TemplateType.proposal.value
        and body.commercialTerms is not None
        and not body.commercialTerms.strip()
    ):
        raise HTTPException(
            status_code=422,
            detail="Proposal templates must include commercial terms.",
        )
    return with_defaults(await templates_repo.update(template_id, body.model_dump(mode="json", exclude_none=True)))


# ── Delete ────────────────────────────────────────────────────────────────────
@router.delete("/{template_id}", status_code=204)
async def delete_template(
    template_id: str,
    user:        CurrentUser = Depends(get_current_user),
):
    template = await templates_repo.get(template_id)
    if not template:
        raise HTTPException(status_code=404, detail="Template not found.")
    if template.get("userId") != user.sub:
        raise HTTPException(status_code=403, detail="Only the owner can delete this template.")
    await templates_repo.delete(template_id)


# ── Apply template (fill in variables) ───────────────────────────────────────
@router.post("/{template_id}/apply", response_model=ApplyTemplateResponse)
async def apply_template(
    template_id: str,
    body:        ApplyTemplateRequest,
    user:        CurrentUser = Depends(get_current_user),
):
    """
    Substitute {variable} placeholders in the template content and commercial terms.
    Standard variables are filled from `proposalId` / `customerId`; `variables` overrides them.
    Placeholders left without a value stay in the text and are listed in `missingVariables`.
    """
    template = await _get_template_for_user(template_id, user)
    values = await _variables_from_sources(body, template, user)
    rendered, missing = render_placeholders(template["content"], values)
    terms, missing_terms = render_placeholders(template.get("commercialTerms"), values)
    # Increment usage counter
    await templates_repo.update(template_id, {"usageCount": template.get("usageCount", 0) + 1})
    return ApplyTemplateResponse(
        templateId=template_id,
        rendered=rendered,
        commercialTerms=terms,
        serviceType=template["serviceType"],
        industry=template["industry"],
        missingVariables=missing + [name for name in missing_terms if name not in missing],
    )
