"""
Extracción de texto de documentos cargados por el usuario (PDF, Word, CSV, Excel).

Usado por el RAG de propuestas (RF-VTA-01): el texto extraído se fragmenta y se
indexa para que la generación use las propuestas históricas de la empresa.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".xlsx", ".txt", ".md"}
UNSUPPORTED_LEGACY = {".doc": ".docx", ".xls": ".xlsx"}

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 400_000
_ROWS_PER_BLOCK = 20


class DocumentExtractionError(ValueError):
    """El archivo no se pudo leer o no contiene texto utilizable."""


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _rows_to_text(header: list[str], rows: list[list[str]], sheet: str | None = None) -> str:
    """Convierte filas tabulares a bloques 'Columna: valor' para que cada fragmento sea autocontenido."""
    header = [h.strip() or f"Columna {i + 1}" for i, h in enumerate(header)]
    lines: list[str] = []
    if sheet:
        lines.append(f"## Hoja: {sheet}")
    for n, row in enumerate(rows, start=1):
        cells = [
            f"{header[i] if i < len(header) else f'Columna {i + 1}'}: {v.strip()}"
            for i, v in enumerate(row)
            if v is not None and str(v).strip()
        ]
        if cells:
            lines.append(" | ".join(cells))
        if n % _ROWS_PER_BLOCK == 0:
            lines.append("")  # corte de párrafo para el chunker
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentExtractionError("El PDF está protegido con contraseña.")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"No se pudo leer el PDF: {exc}") from exc


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentExtractionError(f"No se pudo leer el documento Word: {exc}") from exc

    parts: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style is not None and p.style.name.lower().startswith(("heading", "título", "titulo")):
            parts.append(f"## {text}")
        else:
            parts.append(text)
    for table in doc.tables:
        rows = [[c.text for c in r.cells] for r in table.rows]
        if len(rows) > 1:
            parts.append(_rows_to_text(rows[0], rows[1:]))
        elif rows:
            parts.append(" | ".join(c for c in rows[0] if c.strip()))
    return "\n\n".join(parts)


def _extract_csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if not rows:
        return ""
    if len(rows) == 1:
        return " | ".join(rows[0])
    return _rows_to_text(rows[0], rows[1:])


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentExtractionError(f"No se pudo leer el Excel: {exc}") from exc

    parts: list[str] = []
    for ws in wb.worksheets:
        rows = [
            ["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)
        ]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        parts.append(_rows_to_text(rows[0], rows[1:], sheet=ws.title) if len(rows) > 1
                     else f"## Hoja: {ws.title}\n" + " | ".join(rows[0]))
    wb.close()
    return "\n\n".join(parts)


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".csv": _extract_csv,
    ".xlsx": _extract_xlsx,
    ".txt": lambda d: _decode(d),
    ".md": lambda d: _decode(d),
}


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Devuelve (texto, extensión). Lanza DocumentExtractionError si no es válido."""
    ext = Path(filename or "").suffix.lower()
    if ext in UNSUPPORTED_LEGACY:
        raise DocumentExtractionError(
            f"El formato {ext} no está soportado; guarda el archivo como {UNSUPPORTED_LEGACY[ext]}."
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            "Formato no soportado. Usa: " + ", ".join(sorted(SUPPORTED_EXTENSIONS)) + "."
        )
    if not data:
        raise DocumentExtractionError("El archivo está vacío.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentExtractionError(
            f"El archivo supera el máximo de {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )

    text = _EXTRACTORS[ext](data).replace("\x00", "").strip()
    if not text:
        raise DocumentExtractionError(
            "No se encontró texto en el archivo (¿es un PDF escaneado sin OCR?)."
        )
    return text[:MAX_TEXT_CHARS], ext
